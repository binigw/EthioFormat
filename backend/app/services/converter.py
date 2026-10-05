import os
import gc
import subprocess
import base64
import shutil
from pathlib import Path
from typing import List, Tuple, Dict, Any, Optional
import pymupdf
import docx
from docx.text.paragraph import Paragraph as DocxParagraph
from docx.table import Table as DocxTable
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT

from app.services.pricing_engine import PricingEngine
from app.models.formatting import PricingDetailModel

class DocumentConverterService:
    def __init__(self):
        self.libreoffice_bin = self._find_libreoffice()

    def _find_libreoffice(self) -> str | None:
        candidates = ["libreoffice", "soffice", "/usr/bin/libreoffice", "/usr/bin/soffice"]
        for cand in candidates:
            if shutil.which(cand):
                return cand
        return None

    def _convert_docx_to_pdf_reportlab(self, docx_path: str, pdf_path: str) -> str:
        """
        Pure-Python high-fidelity DOCX to PDF converter using ReportLab & python-docx.
        Applied when LibreOffice is not available in the hosting environment.
        Features intelligent table unwrapping to prevent ReportLab LayoutError on large text cells.
        Adheres strictly to Ethiopian Thesis formatting standards:
        - Margins: 1.5 in Left (binding), 1.0 in Right, Top, Bottom
        - Line spacing: 1.5
        - Headings: Bold, Title Case / Upper Case
        - Justified body text
        """
        doc = docx.Document(docx_path)
        # Printable width for A4 (595.27 pt width - 2.5 in margins)
        printable_width = 595.27 - (1.5 + 1.0) * 72.0

        pdf_doc = SimpleDocTemplate(
            pdf_path,
            pagesize=A4,
            leftMargin=1.5 * inch,
            rightMargin=1.0 * inch,
            topMargin=1.0 * inch,
            bottomMargin=1.0 * inch
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'ThesisTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=15,
            leading=20,
            alignment=TA_CENTER,
            spaceAfter=12
        )

        h1_style = ParagraphStyle(
            'ThesisH1',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=13,
            leading=18,
            spaceBefore=14,
            spaceAfter=8
        )

        h2_style = ParagraphStyle(
            'ThesisH2',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=11.5,
            leading=16,
            spaceBefore=10,
            spaceAfter=6
        )

        body_style = ParagraphStyle(
            'ThesisBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=11,
            leading=17, # 1.5 line spacing
            alignment=TA_JUSTIFY,
            spaceAfter=8
        )

        center_style = ParagraphStyle(
            'ThesisCenter',
            parent=body_style,
            alignment=TA_CENTER
        )

        table_cell_style = ParagraphStyle(
            'TableCell',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13,
            alignment=TA_LEFT
        )

        table_header_style = ParagraphStyle(
            'TableHeader',
            parent=table_cell_style,
            fontName='Helvetica-Bold'
        )

        def format_single_paragraph(p: DocxParagraph) -> Optional[Any]:
            text = p.text.strip()
            if not text:
                return Spacer(1, 6)

            style_name = (p.style.name or '').lower() if p.style else ''
            align_str = str(p.alignment) if p.alignment else ''

            is_title = any(kw in text.upper() for kw in [
                'ADDIS ABABA UNIVERSITY', 'JIMMA UNIVERSITY', 'HAWASSA UNIVERSITY',
                'MEKELLE UNIVERSITY', 'BAHIR DAR UNIVERSITY', 'HARAMAYA UNIVERSITY',
                'A THESIS SUBMITTED', 'IN PARTIAL FULFILLMENT', 'SCHOOL OF GRADUATE STUDIES'
            ])

            is_h1 = (
                'heading 1' in style_name or
                text.startswith('CHAPTER') or
                text in ['TABLE OF CONTENTS', 'ABSTRACT', 'DECLARATION', 'DEDICATION', 'ACKNOWLEDGEMENTS', 'LIST OF TABLES', 'LIST OF FIGURES', 'REFERENCES', 'APPENDIX']
            )

            is_h2 = (
                'heading 2' in style_name or
                ('.' in text[:4] and any(text.startswith(f'{i}.') for i in range(1, 10)))
            )

            formatted_text = ''
            for run in p.runs:
                r_text = run.text
                if not r_text:
                    continue
                r_text = r_text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                if run.bold and run.italic:
                    formatted_text += f'<b><i>{r_text}</i></b>'
                elif run.bold:
                    formatted_text += f'<b>{r_text}</b>'
                elif run.italic:
                    formatted_text += f'<i>{r_text}</i>'
                else:
                    formatted_text += r_text

            if not formatted_text:
                formatted_text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

            if is_title:
                return Paragraph(f'<b>{formatted_text}</b>', title_style)
            elif is_h1:
                return Paragraph(f'<b>{formatted_text}</b>', h1_style)
            elif is_h2:
                return Paragraph(f'<b>{formatted_text}</b>', h2_style)
            elif 'CENTER' in align_str:
                return Paragraph(formatted_text, center_style)
            else:
                return Paragraph(formatted_text, body_style)

        def should_unwrap_table(table: DocxTable) -> bool:
            """
            Checks if a table is used purely for visual layout or contains extensive text
            that would exceed the ReportLab page frame height (~650 points), causing LayoutError.
            """
            if not table.rows or not table.columns:
                return True
            total_cells = len(table.rows) * len(table.columns)
            if total_cells <= 2:
                return True
            for row in table.rows:
                for cell in row.cells:
                    cell_text_len = sum(len(p.text) for p in cell.paragraphs)
                    if cell_text_len > 350 or len(cell.paragraphs) > 3:
                        return True
            return False

        story = []

        # Iterate elements in natural document order
        for child in doc.element.body:
            tag = child.tag.split('}')[-1]
            if tag == 'p':
                p = DocxParagraph(child, doc)
                flowable = format_single_paragraph(p)
                if flowable:
                    story.append(flowable)
            elif tag == 'tbl':
                tbl = DocxTable(child, doc)
                if should_unwrap_table(tbl):
                    # Unwrap table cells into main story flowables
                    for row in tbl.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                flowable = format_single_paragraph(p)
                                if flowable:
                                    story.append(flowable)
                else:
                    # Render standard compact data table with column widths and pagination
                    num_cols = len(tbl.columns)
                    col_width = printable_width / max(1, num_cols)
                    table_data = []
                    for row_idx, row in enumerate(tbl.rows):
                        row_data = []
                        for cell in row.cells:
                            cell_p_list = []
                            for cp in cell.paragraphs:
                                ctext = cp.text.strip().replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                                if ctext:
                                    st = table_header_style if row_idx == 0 else table_cell_style
                                    cell_p_list.append(Paragraph(ctext, st))
                            if not cell_p_list:
                                cell_p_list = [Paragraph('', table_cell_style)]
                            row_data.append(cell_p_list)
                        table_data.append(row_data)

                    if table_data:
                        t = Table(table_data, colWidths=[col_width] * num_cols, splitByRow=1, repeatRows=1)
                        t.setStyle(TableStyle([
                            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F3F4F6')),
                            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#D1D5DB')),
                            ('TOPPADDING', (0,0), (-1,-1), 4),
                            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                        ]))
                        story.append(Spacer(1, 6))
                        story.append(t)
                        story.append(Spacer(1, 8))

        pdf_doc.build(story)
        del doc
        del story
        del pdf_doc
        gc.collect()
        return pdf_path

    def convert_docx_to_pdf(self, docx_path: str, output_dir: str) -> str:
        """
        Converts .docx to .pdf using LibreOffice if available, or ReportLab fallback.
        """
        docx_path_obj = Path(docx_path)
        output_dir_obj = Path(output_dir)
        output_dir_obj.mkdir(parents=True, exist_ok=True)
        pdf_path = output_dir_obj / f"{docx_path_obj.stem}.pdf"

        if self.libreoffice_bin:
            try:
                cmd = [
                    self.libreoffice_bin,
                    "--headless",
                    "--invisible",
                    "--nodefault",
                    "--nofirststartwizard",
                    "--nolockcheck",
                    "--nologo",
                    "--norestore",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(output_dir_obj.resolve()),
                    str(docx_path_obj.resolve())
                ]
                result = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=90)
                if result.returncode == 0 and pdf_path.exists():
                    return str(pdf_path.resolve())
            except Exception as e:
                print(f"[Converter] LibreOffice failed ({e}), falling back to ReportLab...")

        # Fallback to pure Python ReportLab converter
        return self._convert_docx_to_pdf_reportlab(str(docx_path_obj.resolve()), str(pdf_path.resolve()))

    def generate_previews_and_page_count(self, pdf_path: str, max_preview_pages: int = 3) -> Tuple[int, List[str]]:
        """
        Inspects the PDF using PyMuPDF:
        1. Dynamically calculates total page count.
        2. Renders ONLY the first `max_preview_pages` (Pages 1, 2, 3) as high-res base64 PNG images.
        3. Strictly DOES NOT expose or extract pages beyond index (max_preview_pages - 1).
        4. Memory optimized: Explicitly closes doc, deletes pixmaps, and invokes gc.collect().
        """
        doc = pymupdf.open(pdf_path)
        try:
            total_pages = len(doc)

            if total_pages == 0:
                raise ValueError("The generated document has 0 pages.")

            preview_images_base64: List[str] = []
            pages_to_render = min(max_preview_pages, total_pages)

            # 1.5x scale matrix for crisp 108 DPI rendering with minimal RAM footprint
            matrix = pymupdf.Matrix(1.5, 1.5)

            for page_idx in range(pages_to_render):
                page = doc.load_page(page_idx)
                pix = page.get_pixmap(matrix=matrix, alpha=False)
                png_bytes = pix.tobytes(output="png")
                b64_str = base64.b64encode(png_bytes).decode("utf-8")
                preview_images_base64.append(f"data:image/png;base64,{b64_str}")
                
                # Explicit cleanup of PyMuPDF C objects per page
                del pix
                del page
                del png_bytes
        finally:
            doc.close()
            del doc
            gc.collect()

        return total_pages, preview_images_base64

    def process_document_pipeline(self, formatted_docx_path: str, work_dir: str) -> Dict[str, Any]:
        """
        Full pipeline:
        1. DOCX -> PDF conversion
        2. Page count & 3-page base64 preview rendering
        3. Dynamic pricing calculation
        """
        pdf_path = self.convert_docx_to_pdf(formatted_docx_path, work_dir)
        total_pages, preview_pages = self.generate_previews_and_page_count(pdf_path, max_preview_pages=3)
        pricing = PricingEngine.calculate_pricing(total_pages)

        # Trigger garbage collection after rendering
        gc.collect()

        return {
            "total_pages": total_pages,
            "preview_pages": preview_pages,
            "pricing": pricing,
            "pdf_path": pdf_path
        }

converter_service = DocumentConverterService()
