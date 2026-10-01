import docx
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_sample_thesis(filename="sample_ethiopian_thesis.docx"):
    doc = Document()
    
    # 1. Cover Page
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_u = p_univ.add_run("ADDIS ABABA UNIVERSITY\nCOLLEGE OF NATURAL AND COMPUTATIONAL SCIENCES\nDEPARTMENT OF COMPUTER SCIENCE")
    r_u.bold = True
    r_u.font.size = Pt(14)

    for _ in range(3):
        doc.add_paragraph()

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("DESIGN AND IMPLEMENTATION OF MACHINE LEARNING MODELS FOR ETHIOPIC TEXT RECOGNITION AND AUTOMATED CITATION INDEXING")
    r_t.bold = True
    r_t.font.size = Pt(16)

    for _ in range(3):
        doc.add_paragraph()

    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_a = p_author.add_run("BY\nABEBE BIKILA DESTA\n(ID: GSR/1984/14)")
    r_a.bold = True
    r_a.font.size = Pt(12)

    for _ in range(4):
        doc.add_paragraph()

    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_f = p_footer.add_run("A THESIS SUBMITTED TO THE SCHOOL OF GRADUATE STUDIES OF ADDIS ABABA UNIVERSITY IN PARTIAL FULFILLMENT OF THE REQUIREMENTS FOR THE DEGREE OF MASTER OF SCIENCE IN COMPUTER SCIENCE\n\nADDIS ABABA, ETHIOPIA\nJUNE 2026")
    r_f.font.size = Pt(11)

    # Page Break to Page 2
    doc.add_page_break()

    # 2. Approval Sheet
    p_app_h = doc.add_paragraph()
    p_app_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ah = p_app_h.add_run("ADDIS ABABA UNIVERSITY\nSCHOOL OF GRADUATE STUDIES\n\nAPPROVAL SHEET")
    r_ah.bold = True
    r_ah.font.size = Pt(14)

    doc.add_paragraph()
    p_app_body = doc.add_paragraph(
        "This is to certify that the thesis prepared by Abebe Bikila Desta, entitled "
        "'Design and Implementation of Machine Learning Models for Ethiopic Text Recognition and Automated Citation Indexing' "
        "and submitted in partial fulfillment of the requirements for the degree of Master of Science in Computer Science "
        "complies with the regulations of the University and meets the accepted standards with respect to originality and quality."
    )

    doc.add_paragraph()
    p_board = doc.add_paragraph("Signed by the Examining Committee:")
    p_board.runs[0].bold = True

    table = doc.add_table(rows=5, cols=3)
    headers = ["Role", "Name", "Signature & Date"]
    for i, h in enumerate(headers):
        table.cell(0, i).paragraphs[0].text = h
        table.cell(0, i).paragraphs[0].runs[0].bold = True

    roles = ["Advisor", "Co-Advisor", "Internal Examiner", "External Examiner"]
    names = ["Dr. Solomon Gizaw", "Dr. Worku Biratu", "Prof. Mesfin Kebede", "Dr. Almaz Tefera"]
    for idx, (role, name) in enumerate(zip(roles, names), 1):
        table.cell(idx, 0).paragraphs[0].text = role
        table.cell(idx, 1).paragraphs[0].text = name
        table.cell(idx, 2).paragraphs[0].text = "___________________"

    # Page Break to Page 3
    doc.add_page_break()

    # 3. Table of Contents
    p_toc_h = doc.add_paragraph()
    p_toc_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_th = p_toc_h.add_run("TABLE OF CONTENTS")
    r_th.bold = True
    r_th.font.size = Pt(14)

    toc_entries = [
        ("Approval Sheet", "ii"),
        ("Dedication", "iii"),
        ("Acknowledgments", "iv"),
        ("Abstract", "v"),
        ("List of Figures", "vi"),
        ("List of Tables", "vii"),
        ("Acronyms and Abbreviations", "viii"),
        ("CHAPTER 1: INTRODUCTION", "1"),
        ("1.1 Background of the Study", "1"),
        ("1.2 Statement of the Problem", "4"),
        ("1.3 Objectives of the Study", "6"),
        ("1.3.1 General Objective", "6"),
        ("1.3.2 Specific Objectives", "6"),
        ("1.4 Scope and Limitations", "7"),
        ("1.5 Significance of the Study", "8"),
        ("1.6 Thesis Organization", "9"),
        ("CHAPTER 2: LITERATURE REVIEW", "10"),
        ("2.1 Ethiopic Script Characteristics", "10"),
        ("2.2 Deep Learning in OCR Systems", "15"),
        ("2.3 Transformer Architectures in NLP", "22"),
        ("CHAPTER 3: METHODOLOGY", "30"),
        ("3.1 Proposed Architecture", "30"),
        ("3.2 Dataset Collection and Preprocessing", "34"),
        ("3.3 Model Training and Optimization", "42"),
        ("CHAPTER 4: EXPERIMENTAL RESULTS", "50"),
        ("4.1 Evaluation Metrics", "50"),
        ("4.2 Comparative Performance", "55"),
        ("CHAPTER 5: CONCLUSION AND FUTURE WORK", "65"),
        ("5.1 Summary of Findings", "65"),
        ("5.2 Recommendations for Future Research", "68"),
        ("REFERENCES", "71"),
        ("APPENDIX A: SAMPLE DATASETS", "78"),
    ]

    for title, pg in toc_entries:
        p_row = doc.add_paragraph()
        p_row.paragraph_format.space_after = Pt(2)
        dots = "." * (60 - len(title))
        p_row.add_run(f"{title} {dots} {pg}")

    # Page Break to Chapter 1
    doc.add_page_break()

    # 4. Chapter 1 Content
    p_ch1 = doc.add_paragraph()
    p_ch1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c1 = p_ch1.add_run("CHAPTER 1\nINTRODUCTION")
    r_c1.bold = True
    r_c1.font.size = Pt(14)

    p_h11 = doc.add_paragraph()
    r_h11 = p_h11.add_run("1.1 Background of the Study")
    r_h11.bold = True
    r_h11.font.size = Pt(12)

    doc.add_paragraph(
        "Optical Character Recognition (OCR) for indigenous African writing systems remains a critical frontier in computational linguistics and artificial intelligence. "
        "The Ethiopic script (Fidel), with over three hundred syllabic characters and ancient ligatures, presents unique optical disambiguation challenges that standard Latin OCR algorithms fail to address effectively."
    )
    doc.add_paragraph(
        "In recent years, universities across Ethiopia—including Addis Ababa University, Jimma University, Hawassa University, and Bahir Dar University—have accumulated vast repositories of historical manuscripts, legal archives, and postgraduate theses. "
        "Automating the digitization and indexation of these resources requires robust machine learning models capable of handling distinct typographic variations and noisy historical scans."
    )

    p_h12 = doc.add_paragraph()
    r_h12 = p_h12.add_run("1.2 Statement of the Problem")
    r_h12.bold = True
    r_h12.font.size = Pt(12)

    doc.add_paragraph(
        "Existing automated document processing tools consistently misalign Ethiopic font weights, line spacing, and margin requirements specified by the School of Graduate Studies. "
        "Students and researchers waste countless hours manually adjusting tab stops, dot leaders in tables of contents, and binding margins."
    )

    doc.save(filename)
    print(f"Created sample thesis: {filename}")

if __name__ == "__main__":
    create_sample_thesis()
